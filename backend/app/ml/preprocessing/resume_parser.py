"""Resume text extraction (pypdf / python-docx) and light section parsing.

The goal is reliable plain-text extraction plus best-effort identification of
education and experience. Resume layouts vary widely, so the heuristics are
intentionally simple and never block processing when they find nothing.
"""
import io
import re
from datetime import date

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class ResumeParseError(Exception):
    pass


def extract_text_from_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ResumeParseError("The PDF is password protected") from exc
        pages = [(page.extract_text() or "") for page in reader.pages[:30]]
    except ResumeParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ResumeParseError("The PDF file could not be read") from exc
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several exception types for corrupt files
        raise ResumeParseError("The DOCX file could not be read") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text not in cells:
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        text = extract_text_from_pdf(data)
    elif file_type == "docx":
        text = extract_text_from_docx(data)
    else:
        raise ResumeParseError("Unsupported file type")
    text = re.sub(r"[ \t]+", " ", text.replace("\x00", ""))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ---------------------------------------------------------------- sections

SECTION_HEADINGS = {
    "education": ["education", "academic background", "academic qualifications", "qualifications",
                  "educational qualifications", "academics", "academic details", "education and training"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "work history", "internships", "internship", "internship experience", "employment",
                   "career history"],
    "skills": ["skills", "technical skills", "key skills", "core competencies", "technologies",
               "skills summary", "technical expertise", "tools and technologies"],
    "projects": ["projects", "academic projects", "personal projects", "key projects"],
    "certifications": ["certifications", "certificates", "courses", "training"],
    "summary": ["summary", "profile", "professional summary", "objective", "career objective", "about me"],
}
_HEADING_LOOKUP = {h: section for section, heads in SECTION_HEADINGS.items() for h in heads}

DEGREE_RE = re.compile(
    r"\b(b\.?\s?tech|m\.?\s?tech|b\.?\s?e\b|m\.?\s?e\b|b\.?\s?sc|m\.?\s?sc|bca|mca|mba|bba|b\.?\s?com|m\.?\s?com|"
    r"ph\.?\s?d|bachelor|master|diploma|degree|higher secondary|hsc|sslc|ssc|12th|10th|b\.?\s?a\b|m\.?\s?a\b|"
    r"university|college|institute of technology|cgpa|gpa)\b", re.IGNORECASE)
YEARS_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\b(?:\s+of)?(?:\s+\w+){0,3}\s+experience",
                      re.IGNORECASE)
YEARS_SIMPLE_RE = re.compile(r"experience\s*(?:of|:)?\s*(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)", re.IGNORECASE)
MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
RANGE_RE = re.compile(
    rf"(?:(?P<m1>{MONTHS})[a-z]*\.?\s*)?(?P<y1>(?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*"
    rf"(?:(?:(?P<m2>{MONTHS})[a-z]*\.?\s*)?(?P<y2>(?:19|20)\d{{2}})|(?P<present>present|current|now|till date|date))",
    re.IGNORECASE)
_MONTH_INDEX = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun",
                                                "jul", "aug", "sep", "oct", "nov", "dec"])}


def _heading_of(line):
    clean = re.sub(r"[^a-z &]", "", line.lower()).strip()
    if not clean or len(clean) > 40:
        return None
    return _HEADING_LOOKUP.get(clean)


def split_sections(text):
    sections = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = _heading_of(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _month(value):
    return _MONTH_INDEX.get((value or "")[:3].lower(), 1)


def estimate_experience_years(text, experience_lines):
    """Explicit "N years of experience" wins; otherwise sum date ranges in the experience section."""
    explicit = [float(m) for m in YEARS_RE.findall(text) + YEARS_SIMPLE_RE.findall(text)]
    explicit = [v for v in explicit if 0 < v <= 45]
    if explicit:
        return round(max(explicit), 1)

    months = 0
    today = date.today()
    for line in experience_lines:
        for m in RANGE_RE.finditer(line):
            y1 = int(m.group("y1"))
            start = y1 * 12 + _month(m.group("m1"))
            if m.group("present"):
                end = today.year * 12 + today.month
            else:
                end = int(m.group("y2")) * 12 + (_month(m.group("m2")) if m.group("m2") else 12)
            span = end - start + 1  # resume ranges are inclusive: "Jan 2019 - Dec 2020" = 24 months
            if 0 < span <= 45 * 12:
                months += span
    return round(months / 12, 1) if months else None


def parse_sections(text):
    """Return {'education': [...], 'experience': [...], 'experience_years': float|None}."""
    sections = split_sections(text)

    education = [l for l in sections.get("education", []) if len(l) > 3][:6]
    if not education:
        education = [l for l in text.splitlines() if DEGREE_RE.search(l) and 5 < len(l.strip()) < 200][:5]
    education = [l.strip()[:200] for l in education]

    exp_lines = sections.get("experience", [])
    experience = [l.strip()[:200] for l in exp_lines if len(l.strip()) > 3][:10]
    years = estimate_experience_years(text, exp_lines)

    return {"education": education, "experience": experience, "experience_years": years}
